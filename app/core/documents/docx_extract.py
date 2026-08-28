"""DOCX text extraction (Phase 7). python-docx is lazy-imported (it needs
lxml, which some locked-down environments block) — `docx_available()`
returns False on any import failure and the import endpoint degrades to a
503.

DOCX is a flow format: paragraphs and table cells have no page geometry, so
every ExtractedShape here has x/y/w/h = None and is reviewed as a bilingual
reader rather than a positioned canvas. `page_index` is always 0.
"""

import io
from typing import List

from app.core.documents import ExtractedShape

__all__ = ["extract_docx", "docx_available"]


def docx_available() -> bool:
    try:
        import docx  # noqa: F401
    except Exception:  # ImportError, or an OS/DLL-policy block on lxml
        return False
    return True


def _heading_kind(style_name: str) -> str:
    s = (style_name or "").lower()
    if s.startswith("title") or s.startswith("heading"):
        return "title"
    if "caption" in s:
        return "caption"
    return "body"


def extract_docx(raw: bytes) -> List[ExtractedShape]:
    if not docx_available():
        raise RuntimeError("python-docx is not importable in this environment.")

    import docx

    document = docx.Document(io.BytesIO(raw))
    out: List[ExtractedShape] = []
    order = 0

    for para in document.paragraphs:
        text = (para.text or "").strip()
        if not text:
            continue
        out.append(ExtractedShape(
            page_index=0, page_width=0, page_height=0,
            shape_index=order, reading_order=order,
            kind=_heading_kind(getattr(para.style, "name", "")), text=text,
        ))
        order += 1

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                text = (cell.text or "").strip()
                if text:
                    out.append(ExtractedShape(
                        page_index=0, page_width=0, page_height=0,
                        shape_index=order, reading_order=order,
                        kind="table_cell", text=text,
                    ))
                    order += 1

    return out
